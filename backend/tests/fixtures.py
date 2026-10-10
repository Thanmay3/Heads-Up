"""Deterministic test fixtures for material ingestion tests.

make_pdf builds a minimal but valid multi-page PDF from scratch (pypdf cannot
write text). make_pptx uses python-pptx itself, which is deterministic for
fixed content.
"""

import io

from pptx import Presentation


def _escape(text: str) -> str:
    return text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")


def make_pdf(page_texts: list[str]) -> bytes:
    """Build a valid PDF with one text page per entry (\\n = new line)."""
    buf = io.BytesIO()
    buf.write(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}

    def add(obj_num: int, body: bytes) -> None:
        offsets[obj_num] = buf.tell()
        buf.write(f"{obj_num} 0 obj\n".encode("latin-1") + body + b"\nendobj\n")

    font_num = 3
    page_nums = [4 + 2 * i for i in range(len(page_texts))]
    content_nums = [5 + 2 * i for i in range(len(page_texts))]

    add(1, b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{n} 0 R" for n in page_nums)
    add(2, f"<< /Type /Pages /Kids [{kids}] /Count {len(page_texts)} >>".encode("latin-1"))
    add(font_num, b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    for i, page_text in enumerate(page_texts):
        add(
            page_nums[i],
            (
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                f"/Contents {content_nums[i]} 0 R "
                f"/Resources << /Font << /F1 {font_num} 0 R >> >> >>"
            ).encode("latin-1"),
        )
        lines = page_text.split("\n")
        parts = [f"BT /F1 12 Tf 14 TL 72 720 Td ({_escape(lines[0])}) Tj"]
        for line in lines[1:]:
            parts.append(f"T* ({_escape(line)}) Tj")
        parts.append("ET")
        stream = " ".join(parts).encode("latin-1")
        add(
            content_nums[i],
            f"<< /Length {len(stream)} >>\nstream\n".encode("latin-1")
            + stream
            + b"\nendstream",
        )

    max_obj = content_nums[-1]
    xref_pos = buf.tell()
    buf.write(f"xref\n0 {max_obj + 1}\n".encode("latin-1"))
    buf.write(b"0000000000 65535 f \n")
    for num in range(1, max_obj + 1):
        buf.write(f"{offsets[num]:010d} 00000 n \n".encode("latin-1"))
    buf.write(
        f"trailer\n<< /Size {max_obj + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_pos}\n%%EOF\n".encode("latin-1")
    )
    return buf.getvalue()


def make_pptx(slide_lines: list[list[str]]) -> bytes:
    """Build a PPTX; each entry is the list of text lines for one slide."""
    prs = Presentation()
    for lines in slide_lines:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = lines[0]
        slide.placeholders[1].text = "\n".join(lines[1:])
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_blank_pptx(slide_count: int) -> bytes:
    """A PPTX with slides but no text at all (empty-extraction case)."""
    prs = Presentation()
    for _ in range(slide_count):
        prs.slides.add_slide(prs.slide_layouts[6])
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
