"""PPTX text extraction via python-pptx."""

import io
import zipfile

from pptx import Presentation
from pptx.exc import PackageNotFoundError, PythonPptxError

from app.extractors.base import ExtractionError, RawSegment


def _slide_text(slide) -> list[str]:
    lines: list[str] = []
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for paragraph in shape.text_frame.paragraphs:
            text = "".join(run.text for run in paragraph.runs).strip()
            if text:
                lines.append(text)
    return lines


def extract_pptx(payload: bytes) -> list[RawSegment]:
    try:
        presentation = Presentation(io.BytesIO(payload))
    except (PackageNotFoundError, PythonPptxError, zipfile.BadZipFile, ValueError) as exc:
        raise ExtractionError(
            "This file is not a readable PPTX. If it is an old .ppt file, "
            "export it as .pptx first"
        ) from exc

    segments: list[RawSegment] = []
    for index, slide in enumerate(presentation.slides, start=1):
        text = "\n".join(_slide_text(slide))
        if text.strip():
            segments.append(RawSegment(page_number=index, text=text))

    if not segments:
        raise ExtractionError("No text could be extracted from this presentation")
    return segments
